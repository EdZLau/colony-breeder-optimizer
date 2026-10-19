"""Read CSV and XLSX uploads into rows suitable for the core colony loader."""

import csv
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from typing import BinaryIO, Dict, List, Optional, Tuple, Union

OLE_SIGNATURE = bytes.fromhex("D0 CF 11 E0 A1 B1 1A E1")
ROW_LIMIT = 20_000
COLUMN_LIMIT = 1_000
UPLOAD_BYTE_LIMIT = 25 * 1024 * 1024
XLSX_EXPANDED_BYTE_LIMIT = 100 * 1024 * 1024


class _UploadValueError(ValueError):
    """An expected upload validation error that should not be rewrapped."""


class UploadWarning(str):
    """A user-facing warning with its source column when one applies."""

    column: Optional[str]

    def __new__(
        cls,
        message: str,
        column: Optional[str] = None,
    ) -> "UploadWarning":
        warning = str.__new__(cls, message)
        warning.column = column
        return warning


class SheetChoiceRequired(ValueError):
    """Raised when a workbook has multiple sheets and none was selected."""

    def __init__(self, sheet_names: List[str]) -> None:
        self.sheet_names = tuple(sheet_names)
        super().__init__(
            "This workbook has multiple sheets; choose one: "
            + ", ".join(sheet_names)
        )


@dataclass
class UploadResult:
    """Text rows and parsing metadata for an uploaded spreadsheet."""

    fieldnames: List[str]
    rows: List[Dict[Optional[str], object]]
    warnings: List[UploadWarning]
    first_data_row: int
    encoding: Optional[str] = None
    delimiter: Optional[str] = None
    sheet_name: Optional[str] = None


def read_upload(
    file: Union[bytes, bytearray, BinaryIO],
    filename: str = "uploaded data",
    *,
    sheet_name: Optional[str] = None,
    header_row: int = 1,
) -> UploadResult:
    """Read CSV or XLSX by content, preserving data rows as text."""
    if header_row < 1:
        raise ValueError("header_row must be 1 or greater.")
    content = _read_bytes(file)
    if not content:
        raise ValueError("The uploaded file is empty.")
    if content.startswith(OLE_SIGNATURE):
        raise ValueError(
            "This is an old .xls or password-protected Excel file. "
            "Save it as a plain .xlsx file and upload it again."
        )
    if content.startswith(b"PK"):
        return _read_xlsx(content, filename, sheet_name, header_row)
    return _read_csv(content, filename, header_row)


def _read_bytes(file: Union[bytes, bytearray, BinaryIO]) -> bytes:
    if isinstance(file, (bytes, bytearray)):
        content = bytes(file)
    else:
        chunks = []
        total_size = 0
        while total_size <= UPLOAD_BYTE_LIMIT:
            chunk = file.read(UPLOAD_BYTE_LIMIT + 1 - total_size)
            if not isinstance(chunk, bytes):
                raise ValueError("Uploaded file must provide bytes.")
            if not chunk:
                break
            chunks.append(chunk)
            total_size += len(chunk)
        content = b"".join(chunks)
    if len(content) > UPLOAD_BYTE_LIMIT:
        raise ValueError(
            f"Upload exceeds the {UPLOAD_BYTE_LIMIT // (1024 * 1024)} MB file-size limit."
        )
    return content


def _decode_csv(content: bytes) -> Tuple[str, str, List[UploadWarning]]:
    warnings: List[UploadWarning] = []
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return content.decode("utf-16"), "utf-16", warnings
        except UnicodeDecodeError as exc:
            raise ValueError("Could not decode the UTF-16 CSV file.") from exc
    try:
        return content.decode("utf-8-sig"), "utf-8-sig", warnings
    except UnicodeDecodeError:
        try:
            text = content.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise ValueError(
                "CSV is neither valid UTF-8 nor cp1252 text."
            ) from exc
        warnings.append(UploadWarning("CSV was not UTF-8; decoded using cp1252."))
        return text, "cp1252", warnings


def _detect_delimiter(text: str, header_row: int) -> str:
    lines = text.splitlines()[header_row - 1:]
    sample = "\n".join(lines)[:8192]
    counts_by_delimiter: Dict[str, List[int]] = {",": [], ";": [], "\t": []}
    for line in sample.splitlines():
        if not line.strip():
            continue
        for delimiter in counts_by_delimiter:
            try:
                fields = next(
                    csv.reader([line], delimiter=delimiter, strict=True)
                )
                counts_by_delimiter[delimiter].append(len(fields) - 1)
            except csv.Error:
                counts_by_delimiter[delimiter].append(line.count(delimiter))

    def consistency_score(delimiter: str) -> Tuple[int, int, int]:
        counts = counts_by_delimiter[delimiter]
        positive_counts = [count for count in counts if count > 0]
        if not positive_counts:
            return (0, 0, 0)
        frequencies = {
            count: positive_counts.count(count) for count in set(positive_counts)
        }
        return (
            len(positive_counts),
            max(frequencies.values()),
            sum(positive_counts),
        )

    scores = {
        delimiter: consistency_score(delimiter)
        for delimiter in counts_by_delimiter
    }
    best_score = max(scores.values())
    best_delimiters = [
        delimiter for delimiter, score in scores.items() if score == best_score
    ]
    if len(best_delimiters) == 1:
        return best_delimiters[0]
    try:
        return csv.Sniffer().sniff(sample, delimiters="".join(best_delimiters)).delimiter
    except csv.Error:
        return best_delimiters[0]


def _is_empty(value: object) -> bool:
    return value is None or value == ""


def _trim_trailing_empty_rows(rows: List[List[object]]) -> List[List[object]]:
    while rows and all(_is_empty(value) for value in rows[-1]):
        rows.pop()
    return rows


def _row_dicts(
    fieldnames: List[str], records: List[List[object]]
) -> List[Dict[Optional[str], object]]:
    header_width = len(fieldnames)
    result: List[Dict[Optional[str], object]] = []
    for record in records:
        padded = list(record[:header_width])
        if len(padded) < header_width:
            padded.extend([""] * (header_width - len(padded)))
        row: Dict[Optional[str], object] = {
            fieldname: _as_text(value)
            for fieldname, value in zip(fieldnames, padded)
        }
        extra_values = record[header_width:]
        last_extra = next(
            (
                index
                for index in range(len(extra_values) - 1, -1, -1)
                if not _is_empty(extra_values[index])
            ),
            -1,
        )
        if last_extra >= 0:
            row[None] = [
                _as_text(value) for value in extra_values[:last_extra + 1]
            ]
        result.append(row)
    return result


def _as_text(value: object) -> str:
    if _is_empty(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _read_csv(content: bytes, filename: str, header_row: int) -> UploadResult:
    text, encoding, warnings = _decode_csv(content)
    delimiter = _detect_delimiter(text, header_row)
    try:
        records = list(
            csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
        )
    except csv.Error as exc:
        raise ValueError(f"Could not parse CSV file {filename}: {exc}") from exc
    if not records:
        raise ValueError(f"CSV file {filename} is empty.")
    if header_row > len(records):
        raise ValueError(
            f"Header row {header_row} is beyond the end of CSV file {filename}."
        )

    header = records[header_row - 1]
    last_header = max(
        (
            index
            for index in range(len(header))
            if header[index].strip()
            or any(
                index < len(record) and not _is_empty(record[index])
                for record in records[header_row:]
            )
        ),
        default=-1,
    )
    if last_header < 0:
        raise ValueError(f"CSV file {filename} has an empty header row.")
    fieldnames = [value.strip() for value in header[:last_header + 1]]
    if len(fieldnames) > COLUMN_LIMIT:
        raise ValueError(
            f"CSV file exceeds the {COLUMN_LIMIT:,}-column upload limit."
        )
    data_records = _trim_trailing_empty_rows(records[header_row:])
    if len(data_records) > ROW_LIMIT:
        raise ValueError(
            f"CSV file exceeds the {ROW_LIMIT:,}-row upload limit."
        )
    rows = _row_dicts(fieldnames, data_records)
    return UploadResult(
        fieldnames=fieldnames,
        rows=rows,
        warnings=warnings,
        first_data_row=header_row + 1,
        encoding=encoding,
        delimiter=delimiter,
    )


def _read_xlsx(
    content: bytes,
    filename: str,
    sheet_name: Optional[str],
    header_row: int,
) -> UploadResult:
    workbook = None
    try:
        try:
            import openpyxl
        except ImportError as exc:
            raise _UploadValueError(
                "Reading .xlsx files requires openpyxl; install the app dependencies."
            ) from exc

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            expanded_size = sum(member.file_size for member in archive.infolist())
        if expanded_size > XLSX_EXPANDED_BYTE_LIMIT:
            raise _UploadValueError(
                "Excel workbook exceeds the "
                f"{XLSX_EXPANDED_BYTE_LIMIT // (1024 * 1024)} MB expanded-size limit."
            )

        workbook = openpyxl.load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
        sheet_names = list(workbook.sheetnames)
        if not sheet_names:
            raise _UploadValueError(
                f"Excel workbook {filename} contains no worksheets."
            )
        if sheet_name is None:
            if len(sheet_names) > 1:
                raise SheetChoiceRequired(sheet_names)
            sheet_name = sheet_names[0]
        elif sheet_name not in sheet_names:
            raise _UploadValueError(
                f"Worksheet {sheet_name!r} not found in {filename}. Available sheets: "
                + ", ".join(sheet_names)
            )

        worksheet = workbook[sheet_name]
        worksheet_rows = iter(
            worksheet.iter_rows(
                values_only=True,
                max_col=COLUMN_LIMIT + 1,
            )
        )
        header_values = None
        rows_read = 0
        for worksheet_row, values in enumerate(worksheet_rows, start=1):
            rows_read = worksheet_row
            if worksheet_row == header_row:
                header_values = list(values)
                break
        if header_values is None:
            if rows_read == 0:
                raise _UploadValueError(
                    f"Worksheet {sheet_name!r} in {filename} has an empty header row."
                )
            raise _UploadValueError(
                f"Header row {header_row} is beyond the end of worksheet {sheet_name!r}."
            )

        while header_values and _is_empty(header_values[-1]):
            header_values.pop()
        if not header_values:
            raise _UploadValueError(
                f"Worksheet {sheet_name!r} in {filename} has an empty header row."
            )
        if len(header_values) > COLUMN_LIMIT:
            raise _UploadValueError(
                f"Excel worksheet exceeds the {COLUMN_LIMIT:,}-column upload limit."
            )
        fieldnames = [_as_text(value).strip() for value in header_values]

        raw_records: List[List[object]] = []
        pending_blank_rows = 0
        date_counts: Dict[str, List[int]] = {}
        boolean_counts: Dict[str, List[int]] = {}
        warnings: List[UploadWarning] = []

        for worksheet_row, values in enumerate(
            worksheet_rows, start=header_row + 1
        ):
            record = list(values)
            while record and _is_empty(record[-1]):
                record.pop()

            if not any(not _is_empty(value) for value in record):
                pending_blank_rows += 1
                if pending_blank_rows > ROW_LIMIT:
                    warnings.append(UploadWarning(
                        f"Data below worksheet row {worksheet_row} was ignored "
                        f"after {ROW_LIMIT:,} consecutive blank rows."
                    ))
                    break
                continue
            if len(record) > COLUMN_LIMIT:
                raise _UploadValueError(
                    f"Excel worksheet exceeds the {COLUMN_LIMIT:,}-column upload limit."
                )
            if len(raw_records) + pending_blank_rows + 1 > ROW_LIMIT:
                raise _UploadValueError(
                    f"Excel worksheet exceeds the {ROW_LIMIT:,}-row upload limit."
                )
            raw_records.extend([[] for _ in range(pending_blank_rows)])
            pending_blank_rows = 0
            raw_records.append(record)

            for column_index, value in enumerate(record):
                name = (
                    fieldnames[column_index]
                    if column_index < len(fieldnames)
                    else ""
                )
                if isinstance(value, bool):
                    boolean_counts.setdefault(name, []).append(worksheet_row)
                elif isinstance(value, (datetime, date, time, timedelta)):
                    date_counts.setdefault(name, []).append(worksheet_row)

        for column, row_numbers in date_counts.items():
            warnings.append(UploadWarning(
                f"Column {column!r} contains {len(row_numbers)} date/time-typed "
                f"cell(s); first on row {row_numbers[0]}. Dates are not repaired.",
                column=column,
            ))
        for column, row_numbers in boolean_counts.items():
            warnings.append(UploadWarning(
                f"Column {column!r} contains {len(row_numbers)} boolean cell(s); "
                f"first on row {row_numbers[0]}. Values are converted to text.",
                column=column,
            ))

        return UploadResult(
            fieldnames=fieldnames,
            rows=_row_dicts(fieldnames, raw_records),
            warnings=warnings,
            first_data_row=header_row + 1,
            sheet_name=sheet_name,
        )
    except (SheetChoiceRequired, _UploadValueError):
        raise
    except Exception as exc:
        raise ValueError(
            f"Could not read {filename} as a valid .xlsx workbook. "
            "Check that it is not corrupt and save it as a plain .xlsx file."
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()
